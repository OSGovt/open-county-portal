"""Reading exported tables from old systems: CSV, TSV, pipe-delimited text
and Excel, in UTF-8 or old Windows encodings. Read-only, always."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from pathlib import Path
from typing import Iterator

TABLE_EXTS = {".csv", ".tsv", ".txt", ".psv", ".xlsx"}
DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%Y%m%d", "%d-%b-%Y", "%b %d, %Y", "%B %d, %Y",
                "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %I:%M:%S %p", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"]


def detect_encoding(path: Path) -> str:
    raw = path.read_bytes()[:2_000_000]
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        return "cp1252"


def detect_delimiter(path: Path, encoding: str) -> str:
    sample = path.read_text(encoding=encoding, errors="replace")[:50_000]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t|;").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (",", "\t", "|", ";")}
        return max(counts, key=counts.get)


def read_rows(path: Path, encoding: str | None = None, delimiter: str | None = None,
              sheet: str | None = None) -> Iterator[dict[str, str]]:
    """Yield each row as {column: text}. Column names are stripped of spaces."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else f"column_{i + 1}" for i, h in enumerate(next(rows, []))]
        for values in rows:
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            out = {}
            for h, v in zip(header, values):
                if isinstance(v, dt.datetime):
                    v = v.date().isoformat() if v.time() == dt.time() else v.isoformat()
                elif isinstance(v, dt.date):
                    v = v.isoformat()
                elif isinstance(v, float) and v.is_integer():
                    v = str(int(v))
                out[h] = "" if v is None else str(v).strip()
            yield out
        wb.close()
        return
    encoding = encoding or detect_encoding(path)
    delimiter = delimiter or detect_delimiter(path, encoding)
    with path.open(encoding=encoding, errors="replace", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=delimiter)
        if reader.fieldnames:
            reader.fieldnames = [(f or "").strip() for f in reader.fieldnames]
        for row in reader:
            if not any((v or "").strip() for v in row.values() if isinstance(v, str)):
                continue
            yield {k: (v or "").strip() if isinstance(v, str) else "" for k, v in row.items() if k is not None}


def parse_date(text: str, fmt: str | None = None) -> dt.date | None:
    text = (text or "").strip()
    if not text:
        return None
    for f in ([fmt] if fmt else []) + DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, f).date()
        except (ValueError, TypeError):
            continue
    return None


def guess_date_format(samples: list[str]) -> str | None:
    best, best_hits = None, 0
    for f in DATE_FORMATS:
        hits = 0
        for s in samples:
            try:
                dt.datetime.strptime(s.strip(), f)
                hits += 1
            except ValueError:
                pass
        if hits > best_hits:
            best, best_hits = f, hits
    return best if best_hits >= max(1, len(samples) // 2) else None


def profile(path: Path, max_rows: int = 200_000) -> dict:
    """Describe a table: encoding, delimiter, rows, and each column's fill rate and samples."""
    info: dict = {"file": path.name, "bytes": path.stat().st_size}
    if path.suffix.lower() != ".xlsx":
        info["encoding"] = detect_encoding(path)
        info["delimiter"] = detect_delimiter(path, info["encoding"])
    cols: dict[str, dict] = {}
    n = 0
    for row in read_rows(path, info.get("encoding"), info.get("delimiter")):
        n += 1
        for k, v in row.items():
            c = cols.setdefault(k, {"filled": 0, "samples": [], "max_len": 0})
            if v:
                c["filled"] += 1
                c["max_len"] = max(c["max_len"], len(v))
                if len(c["samples"]) < 8 and v not in c["samples"]:
                    c["samples"].append(v[:80])
        if n >= max_rows:
            info["truncated"] = True
            break
    info["rows"] = n
    for k, c in cols.items():
        c["fill_rate"] = round(c["filled"] / n, 3) if n else 0
        c["looks_like_date"] = guess_date_format(c["samples"]) if c["samples"] else None
    info["columns"] = cols
    return info


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())
